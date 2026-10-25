import argparse
import csv
from pathlib import Path

import pandas as pd
import openpyxl

from gates import audit_gates

ROOT = Path(__file__).resolve().parents[1]

DILI_SPLITS = {
    "DILI scaffold": "scaffold",
    "DILI external": "external_dilist",
    "DILI temporal": "temporal",
    "DILI cluster-out": "cluster_out",
}
TOX21_SPLITS = {
    "Tox21 NR-AR scaffold": ("nrar", "scaffold"),
    "Tox21 NR-AR cluster-out": ("nrar", "cluster_out"),
    "Tox21 NR-ER scaffold": ("nrer", "scaffold"),
    "Tox21 NR-ER cluster-out": ("nrer", "cluster_out"),
}
NASA_SPLITS = {
    "NASA KC1": "kc1_internal",
    "NASA PC1": "kc1_pc1",
    "NASA JM1": "kc1_jm1",
    "NASA KC2": "kc1_kc2",
}


def read_s17(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["S17_Gate_Containment"]
    out = {}
    for r in range(4, 16):
        name = ws.cell(r, 1).value
        if not name:
            continue
        out[name] = {
            "th_spec95": float(ws.cell(r, 2).value),
            "q_hat": float(ws.cell(r, 9).value),
            "removed": int(ws.cell(r, 7).value),
        }
    return out


def read_dili_tables(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    u80 = {}
    ws = wb["S10_Uncertainty_AD"]
    for r in range(6, ws.max_row + 1):
        split = ws.cell(r, 1).value
        model = ws.cell(r, 2).value
        mode = ws.cell(r, 3).value
        val = ws.cell(r, 5).value
        if split and model == "HistGB" and mode == "validation_U80" and val is not None:
            u80[split] = float(val)
    sens = {}
    ws = wb["S9_Low_Risk"]
    for r in range(6, ws.max_row + 1):
        split = ws.cell(r, 1).value
        model = ws.cell(r, 2).value
        rule = ws.cell(r, 3).value
        val = ws.cell(r, 4).value
        if split and model == "HistGB" and rule == "th_sens95" and val is not None:
            sens[split] = float(val)
    low = {}
    ws = wb["S8_Triage"]
    for r in range(6, 10):
        low[ws.cell(r, 1).value] = int(ws.cell(r, 7).value)
    return u80, sens, low


def read_tox21_low(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["S12_Tox21"]
    out = {}
    for r in range(6, ws.max_row + 1):
        dataset = ws.cell(r, 1).value
        split = ws.cell(r, 2).value
        if dataset in ("Tox21 NR-AR", "Tox21 NR-ER"):
            key = ("nrar" if dataset == "Tox21 NR-AR" else "nrer", str(split).replace("-", "_"))
            out[key] = int(ws.cell(r, 11).value)
    return out


def read_nasa_low(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["S13_NASA_MDP"]
    out = {}
    for r in range(6, ws.max_row + 1):
        split = ws.cell(r, 2).value
        if split:
            out[split] = int(ws.cell(r, 19).value)
    return out


def tox21_run(root, assay, split, q_hat, u80, th_spec, th_sens):
    base = Path(root) / ("tox21_nrar" if assay == "nrar" else "tox21_nrer")
    cs = pd.read_csv(base / "predictions" / f"conformal_sets_{split}_HistGB_alpha0.1.csv")
    ad = pd.read_csv(base / "predictions" / f"ad_features_{split}.csv")
    m = cs.merge(ad[["sample_id", "ood_extreme_flag"]], on="sample_id", how="left")
    m = m[m["ood_extreme_flag"] != True]
    return audit_gates(m["calibrated_score"].values, q_hat, u80, th_spec, th_sens)


def nasa_run(root, split, q_hat, u80, th_spec, th_sens):
    base = Path(root) / "predictions"
    cs = pd.read_csv(base / f"conformal_sets_{split}_HistGB_alpha0.1.csv")
    ad = pd.read_csv(base / f"ad_features_{split}.csv")
    m = cs.merge(ad[["sample_id", "ood_extreme_flag"]], on="sample_id", how="left")
    m = m[m["ood_extreme_flag"] != True]
    return audit_gates(m["calibrated_score"].values, q_hat, u80, th_spec, th_sens)


def threshold_row(path, split):
    th = pd.read_csv(path)
    row = th[(th["split"] == split) & (th["model"] == "HistGB")].iloc[0]
    return float(row["u80_threshold"]), float(row["th_spec95"]), float(row["th_sens95"])


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--data-s1", default=str(ROOT / "results" / "Data S1.xlsx"))
    ap.add_argument("--dili-pred", default=str(ROOT / "results" / "figure_data" / "dili" / "canonical" / "predictions" / "histgb_predictions.csv"))
    ap.add_argument("--tox21-root", default=str(ROOT / "results" / "predictions" / "cross_domain"))
    ap.add_argument("--nasa-root", default=str(ROOT / "results" / "predictions" / "cross_domain" / "nasa"))
    ap.add_argument("--out", default=str(ROOT / "audit" / "verification.csv"))
    args = ap.parse_args()

    s17 = read_s17(args.data_s1)
    u80, sens, dili_low = read_dili_tables(args.data_s1)
    dili = pd.read_csv(args.dili_pred)

    rows = []
    for name, split in DILI_SPLITS.items():
        p = dili[dili["split"] == split]["calibrated_probability"].values
        r = audit_gates(p, s17[name]["q_hat"], u80[split], s17[name]["th_spec95"], sens[split])
        rows.append([name, r["gate"], r["n_high_removed"], s17[name]["removed"], r["n_low"], dili_low[split]])

    if args.tox21_root:
        tox_low = read_tox21_low(args.data_s1)
        for name, (assay, split) in TOX21_SPLITS.items():
            base = Path(args.tox21_root) / ("tox21_nrar" if assay == "nrar" else "tox21_nrer")
            u, ts, tse = threshold_row(base / "tables" / "threshold_lock_table.csv", split)
            r = tox21_run(args.tox21_root, assay, split, s17[name]["q_hat"], u, ts, tse)
            rows.append([name, r["gate"], r["n_high_removed"], s17[name]["removed"], r["n_low"], tox_low[(assay, split)]])

    if args.nasa_root:
        nasa_low = read_nasa_low(args.data_s1)
        for name, split in NASA_SPLITS.items():
            u, ts, tse = threshold_row(Path(args.nasa_root) / "tables" / "threshold_lock_table.csv", split)
            r = nasa_run(args.nasa_root, split, s17[name]["q_hat"], u, ts, tse)
            rows.append([name, r["gate"], r["n_high_removed"], s17[name]["removed"], r["n_low"], nasa_low[split]])

    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    with out.open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["setting", "gate", "removed", "expected_removed", "low_n", "expected_low_n", "pass"])
        ok = 0
        for row in rows:
            passed = row[2] == row[3] and row[4] == row[5]
            ok += int(passed)
            w.writerow(row + [passed])
    print(f"{ok}/{len(rows)} PASS")
    for row in rows:
        if row[2] != row[3] or row[4] != row[5]:
            print("FAIL", row)


if __name__ == "__main__":
    main()
